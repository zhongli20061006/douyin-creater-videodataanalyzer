import json
import os
import subprocess
import threading
import time as time_module
import csv
import io
import tempfile
from contextlib import asynccontextmanager
import pymysql
import redis
from datetime import datetime
from typing import Optional

from fastapi import FastAPI, Query, HTTPException, Header, Depends
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse, RedirectResponse, Response, StreamingResponse
from pydantic import BaseModel

import quality as quality_service
import collector
import queue_service
import extension_receiver
import analyzer
import cleanup_service
import export_service
from time_filter import build_publish_filter

MYSQL_HOST = 'localhost'
MYSQL_PORT = 3307
MYSQL_USER = 'root'
MYSQL_PASSWORD = ''
MYSQL_DB = 'douyin_spider'
REDIS_HOST = 'localhost'
REDIS_PORT = 6379
REDIS_PARAMS = {}
REDIS_START_URLS_KEY = 'douyin:start_urls'

# 优先从 local_config 读 MySQL 配置（开源版无 Scrapy）；缺失时回退 Scrapy settings。
try:
    from local_config import (
        MYSQL_HOST as _H, MYSQL_PORT as _P, MYSQL_USER as _U,
        MYSQL_PASSWORD as _PW, MYSQL_DB as _DB,
    )
    MYSQL_HOST, MYSQL_PORT, MYSQL_USER, MYSQL_PASSWORD, MYSQL_DB = _H, _P, _U, _PW, _DB
except ImportError:
    try:
        from scrapy.utils.project import get_project_settings
        _s = get_project_settings()
        MYSQL_HOST = _s.get('MYSQL_HOST', MYSQL_HOST)
        MYSQL_PORT = _s.getint('MYSQL_PORT', MYSQL_PORT)
        MYSQL_USER = _s.get('MYSQL_USER', MYSQL_USER)
        MYSQL_PASSWORD = _s.get('MYSQL_PASSWORD', MYSQL_PASSWORD)
        MYSQL_DB = _s.get('MYSQL_DB', MYSQL_DB)
        REDIS_HOST = _s.get('REDIS_HOST', REDIS_HOST)
        REDIS_PORT = _s.getint('REDIS_PORT', REDIS_PORT)
        REDIS_PARAMS = _s.getdict('REDIS_PARAMS', {})
        REDIS_START_URLS_KEY = _s.get('REDIS_START_URLS_KEY', REDIS_START_URLS_KEY)
    except Exception:
        pass

VIDEO_IDS_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'video_ids.txt')
CLEANUP_CONFIG_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'cleanup_config.json')

try:
    from local_config import EXTENSION_API_TOKEN, ALLOWED_AUTHOR_IDS, CLEANUP_STORAGE
except Exception:
    EXTENSION_API_TOKEN = ''
    ALLOWED_AUTHOR_IDS = []
    CLEANUP_STORAGE = 'redis'

ALLOWED_ORIGINS = [
    'http://127.0.0.1:8001',
    'http://localhost:8001',
    'http://localhost:5173',
    'http://47.120.36.73',
]

def _cleanup_once() -> None:
    """执行一次清理检查：满足条件则按作者规则备份并删除。"""
    cfg = _read_cleanup_config()
    enabled = bool(cfg['enabled'])
    last_raw = cfg['last_clean_time']
    last = None
    if last_raw:
        try:
            last = datetime.fromisoformat(last_raw)
        except ValueError:
            last = None
    if not cleanup_service.should_run_cleanup(enabled, last, datetime.now()):
        return
    batch_size = int(cfg['batch_size'])
    authors = list(cfg['authors'])

    db = get_db()
    try:
        with db.cursor() as cursor:
            # 只取分组与排序所需三列，避免全表全字段进内存
            cursor.execute('SELECT video_id, author_id, update_time FROM video_info')
            light_rows = cursor.fetchall()
    finally:
        db_close(db)
    if not light_rows:
        return

    ids = cleanup_service.select_stale_ids_per_author(
        light_rows, batch_size=batch_size, author_ids=authors or None,
    )
    if not ids:
        print('定时清理跳过：没有满足条件的待删数据')
        return

    db = get_db()
    try:
        with db.cursor() as cursor:
            placeholders = ', '.join(['%s'] * len(ids))
            cursor.execute(
                f'SELECT * FROM video_info WHERE video_id IN ({placeholders})',
                tuple(ids),
            )
            delete_rows = cursor.fetchall()
    finally:
        db_close(db)

    backup_dir = cleanup_service.CLEANUP_BACKUP_DIR
    os.makedirs(backup_dir, exist_ok=True)
    backup_path = os.path.join(
        backup_dir,
        'cleanup_' + datetime.now().strftime('%Y%m%d_%H%M%S') + '.csv',
    )
    with open(backup_path, 'w', encoding='utf-8', newline='') as f:
        f.write(cleanup_service.build_backup_csv(delete_rows))

    db = get_db()
    try:
        with db.cursor() as cursor:
            placeholders = ', '.join(['%s'] * len(ids))
            cursor.execute(
                f'DELETE FROM video_info WHERE video_id IN ({placeholders})',
                tuple(ids),
            )
            db.commit()
    finally:
        db_close(db)

    cfg['last_clean_time'] = datetime.now().isoformat(timespec='seconds')
    _write_cleanup_config(cfg)
    print(f'定时清理完成：删除 {len(ids)} 条，备份 {backup_path}')


def _cleanup_loop() -> None:
    """后台循环：每 24 小时检查一次清理条件。"""
    while True:
        try:
            _cleanup_once()
        except Exception as e:  # noqa: BLE001 - 后台线程兜底，避免循环退出
            print(f'定时清理异常：{e}')
        time_module.sleep(24 * 3600)


@asynccontextmanager
async def lifespan(_app):
    """应用启动时注册定时清理后台线程（daemon，进程退出自动结束）。"""
    threading.Thread(target=_cleanup_loop, daemon=True).start()
    yield


app = FastAPI(title='抖音爬虫管理面板', version='1.0.0', lifespan=lifespan)

app.add_middleware(
    CORSMiddleware,
    allow_origins=ALLOWED_ORIGINS,
    allow_credentials=True,
    allow_methods=['*'],
    allow_headers=['*'],
)


def verify_write_guard(
    origin: Optional[str] = Header(default=None),
    x_api_token: Optional[str] = Header(default=None, alias='X-API-Token'),
) -> None:
    """写接口守卫：Origin 白名单或 X-API-Token 通过；未配置令牌时 fail-closed。"""
    allowed, status_code, reason = extension_receiver.evaluate_write_guard(
        origin, x_api_token, EXTENSION_API_TOKEN, ALLOWED_ORIGINS,
    )
    if not allowed:
        raise HTTPException(status_code=status_code, detail=reason)


def verify_read_guard(
    origin: Optional[str] = Header(default=None),
    x_api_token: Optional[str] = Header(default=None, alias='X-API-Token'),
) -> None:
    """只读接口守卫：与写守卫同语义（Origin 白名单或 X-API-Token；未配置令牌时 fail-closed）。

    公网下 GET 接口也强制令牌，避免数据裸奔；本机 Origin 白名单仍放行。
    """
    allowed, status_code, reason = extension_receiver.evaluate_write_guard(
        origin, x_api_token, EXTENSION_API_TOKEN, ALLOWED_ORIGINS,
    )
    if not allowed:
        raise HTTPException(status_code=status_code, detail=reason)


def get_redis():
    return redis.Redis(
        host=REDIS_HOST,
        port=REDIS_PORT,
        password=REDIS_PARAMS.get('password'),
        db=REDIS_PARAMS.get('db', 0),
        decode_responses=True,
    )


def get_db():
    try:
        return pymysql.connect(
            host=MYSQL_HOST,
            port=MYSQL_PORT,
            user=MYSQL_USER,
            password=MYSQL_PASSWORD,
            database=MYSQL_DB,
            charset='utf8mb4',
            cursorclass=pymysql.cursors.DictCursor,
            connect_timeout=5,
        )
    except pymysql.Error as e:
        raise HTTPException(status_code=503, detail=f'MySQL 连接失败: {e}')


def apply_publish_filter(start_date: str, end_date: str):
    """把 start_date/end_date 转成 SQL 过滤条件；非法参数抛 400。"""
    try:
        return build_publish_filter(start_date or None, end_date or None)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


def _check_export_total(total: int) -> None:
    if total > export_service.EXPORT_MAX_ROWS:
        raise HTTPException(status_code=400, detail=f'数据量过大（{total} 条），请缩小筛选范围后导出')


def _read_cleanup_config() -> dict:
    """读取清理配置：本地版默认 Redis，开源版（CLEANUP_STORAGE=json）用 JSON 文件。"""
    if CLEANUP_STORAGE == 'json':
        return cleanup_service.read_cleanup_config(CLEANUP_CONFIG_PATH)
    r = get_redis()
    enabled = int(r.get('douyin:cleanup_enabled') or 0) == 1
    last_clean_time = r.get('douyin:cleanup_last_time')
    batch_size = int(r.get('douyin:cleanup_batch_size') or cleanup_service.CLEANUP_BATCH_SIZE)
    authors_raw = r.get('douyin:cleanup_authors')
    authors = json.loads(authors_raw) if authors_raw else []
    return {
        'enabled': enabled,
        'last_clean_time': last_clean_time,
        'batch_size': batch_size,
        'authors': authors,
    }


def _write_cleanup_config(cfg: dict) -> None:
    """写清理配置：与 _read_cleanup_config 对应的双存储。"""
    if CLEANUP_STORAGE == 'json':
        cleanup_service.write_cleanup_config(CLEANUP_CONFIG_PATH, cfg)
        return
    r = get_redis()
    r.set('douyin:cleanup_enabled', '1' if cfg.get('enabled') else '0')
    if cfg.get('last_clean_time'):
        r.set('douyin:cleanup_last_time', cfg['last_clean_time'])
    r.set('douyin:cleanup_batch_size', str(cfg.get('batch_size', cleanup_service.CLEANUP_BATCH_SIZE)))
    r.set('douyin:cleanup_authors', json.dumps(list(cfg.get('authors', []))))


def db_close(db):
    try:
        db.close()
    except Exception:
        pass


# ── Spider Process Manager ──

class SpiderManager:
    def __init__(self):
        self.process: Optional[subprocess.Popen] = None
        self.started_at: Optional[datetime] = None
        self.project_root = os.path.dirname(os.path.abspath(__file__))
        if os.name == 'nt':
            self.venv_python = os.path.join(self.project_root, '.venv', 'Scripts', 'python.exe')
        else:
            self.venv_python = os.path.join(self.project_root, '.venv', 'bin', 'python')
        self.log_path = os.path.join(self.project_root, 'spider_output.log')

    def start(self):
        if self.is_alive():
            return False, '爬虫已在运行'
        log_file = open(self.log_path, 'a', encoding='utf-8')
        log_file.write(f'\n=== Spider started at {datetime.now().isoformat()} ===\n')
        self.process = subprocess.Popen(
            [self.venv_python, 'start_spider.py', '--mode', 'start'],
            cwd=self.project_root,
            stdout=log_file,
            stderr=subprocess.STDOUT,
        )
        self.started_at = datetime.now()
        return True, '爬虫已启动'

    def stop(self):
        if not self.is_alive():
            return False, '爬虫未运行'
        self.process.terminate()
        try:
            self.process.wait(timeout=10)
        except subprocess.TimeoutExpired:
            self.process.kill()
            self.process.wait()
        self.process = None
        self.started_at = None
        return True, '爬虫已停止'

    def is_alive(self):
        return self.process is not None and self.process.poll() is None

    def get_status(self):
        return {
            'running': self.is_alive(),
            'pid': self.process.pid if self.process else None,
            'started_at': self.started_at.isoformat() if self.started_at else None,
        }

    def get_log(self, lines=50):
        return read_log_tail(self.log_path, lines)


spider_manager = SpiderManager()


def decode_log_bytes(raw):
    """日志字节解码：逐行处理，每行 UTF-8 优先，失败回退 GBK（兼容混合编码文件）。"""
    decoded = []
    for line in raw.splitlines():
        try:
            decoded.append(line.decode('utf-8'))
        except UnicodeDecodeError:
            decoded.append(line.decode('gbk', errors='replace'))
    return '\n'.join(decoded)


def read_log_tail(path, lines=50):
    """读取日志文件末尾 N 行；兼容混合编码。"""
    if not os.path.exists(path):
        return []
    with open(path, 'rb') as f:
        raw = f.read()
    return decode_log_bytes(raw).splitlines()[-lines:]


# ── Pydantic Models ──

class CrawlRequest(BaseModel):
    video_ids: list[str]
    task_type: str = 'video'


class CrawlResponse(BaseModel):
    pushed: int
    queue_length: int
    video_ids: list[str]
    skipped: int = 0


class VideoItem(BaseModel):
    video_id: str
    video_title: Optional[str] = None
    video_desc: Optional[str] = None
    author_name: Optional[str] = None
    author_id: Optional[str] = None
    publish_time: Optional[datetime] = None
    like_count: Optional[int] = None
    comment_count: Optional[int] = None
    share_count: Optional[int] = None
    collect_count: Optional[int] = None
    play_count: Optional[int] = None
    video_url: Optional[str] = None
    cover_url: Optional[str] = None
    crawl_time: Optional[datetime] = None
    update_time: Optional[datetime] = None


class PaginatedResponse(BaseModel):
    total: int
    page: int
    page_size: int
    total_pages: int
    data: list[VideoItem]


class StatsResponse(BaseModel):
    total_videos: int
    total_authors: int
    total_likes: int
    total_comments: int
    total_shares: int
    total_plays: int
    latest_crawl: Optional[datetime] = None
    queue_length: int = 0


# ── API Endpoints ──

@app.get('/api/videos', response_model=PaginatedResponse, dependencies=[Depends(verify_read_guard)])
def list_videos(
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    search: str = Query('', description='搜索视频标题/作者/ID'),
    sort_by: str = Query('crawl_time', description='排序字段'),
    order: str = Query('desc', pattern='^(asc|desc)$'),
    start_date: str = Query('', description='发布时间起始（YYYY-MM-DD）'),
    end_date: str = Query('', description='发布时间结束（YYYY-MM-DD）'),
):
    allowed_sort = {
        'video_id', 'video_title', 'author_name', 'publish_time',
        'like_count', 'comment_count', 'share_count', 'play_count', 'collect_count',
        'crawl_time', 'update_time',
    }
    if sort_by not in allowed_sort:
        sort_by = 'crawl_time'

    order_clause = 'DESC' if order == 'desc' else 'ASC'
    offset = (page - 1) * page_size
    publish_clause, publish_params = apply_publish_filter(start_date, end_date)
    author_clause, author_params = extension_receiver.build_author_filter(ALLOWED_AUTHOR_IDS)

    db = get_db()
    try:
        with db.cursor() as cursor:
            if search:
                search_param = f'%{search}%'
                where_parts = ['(video_id LIKE %s OR video_title LIKE %s OR author_name LIKE %s)']
                count_params = [search_param, search_param, search_param]
                if publish_clause:
                    where_parts.append(publish_clause)
                    count_params.extend(publish_params)
                if author_clause:
                    where_parts.append(author_clause)
                    count_params.extend(author_params)
                where_sql = ' AND '.join(where_parts)
                count_sql = f'SELECT COUNT(*) AS total FROM video_info WHERE {where_sql}'
                cursor.execute(count_sql, tuple(count_params))
                total = cursor.fetchone()['total']

                data_sql = f"""
                    SELECT * FROM video_info
                    WHERE {where_sql}
                    ORDER BY {sort_by} {order_clause}
                    LIMIT %s OFFSET %s
                """
                cursor.execute(data_sql, tuple(count_params + [page_size, offset]))
            else:
                where_parts = []
                count_params = []
                if publish_clause:
                    where_parts.append(publish_clause)
                    count_params.extend(publish_params)
                if author_clause:
                    where_parts.append(author_clause)
                    count_params.extend(author_params)
                where_sql = ('WHERE ' + ' AND '.join(where_parts)) if where_parts else ''
                cursor.execute(f'SELECT COUNT(*) AS total FROM video_info {where_sql}', tuple(count_params))
                total = cursor.fetchone()['total']

                data_sql = f"""
                    SELECT * FROM video_info
                    {where_sql}
                    ORDER BY {sort_by} {order_clause}
                    LIMIT %s OFFSET %s
                """
                cursor.execute(data_sql, tuple(count_params + [page_size, offset]))

            rows = cursor.fetchall()

        return PaginatedResponse(
            total=total,
            page=page,
            page_size=page_size,
            total_pages=(total + page_size - 1) // page_size,
            data=[VideoItem(**row) for row in rows],
        )
    finally:
        db_close(db)


@app.get('/api/videos/{video_id}', response_model=VideoItem, dependencies=[Depends(verify_read_guard)])
def get_video(video_id: str):
    db = get_db()
    try:
        with db.cursor() as cursor:
            cursor.execute('SELECT * FROM video_info WHERE video_id = %s', (video_id,))
            row = cursor.fetchone()
        if not row:
            raise HTTPException(status_code=404, detail='视频不存在')
        return VideoItem(**row)
    finally:
        db_close(db)


@app.get('/api/stats', response_model=StatsResponse, dependencies=[Depends(verify_read_guard)])
def get_stats(
    start_date: str = Query('', description='发布时间起始（YYYY-MM-DD）'),
    end_date: str = Query('', description='发布时间结束（YYYY-MM-DD）'),
):
    db = get_db()
    try:
        with db.cursor() as cursor:
            publish_clause, publish_params = apply_publish_filter(start_date, end_date)
            where_sql = ('WHERE ' + publish_clause) if publish_clause else ''
            cursor.execute(f"""
                SELECT
                    COUNT(*) AS total_videos,
                    COUNT(DISTINCT author_id) AS total_authors,
                    COALESCE(SUM(like_count), 0) AS total_likes,
                    COALESCE(SUM(comment_count), 0) AS total_comments,
                    COALESCE(SUM(share_count), 0) AS total_shares,
                    COALESCE(SUM(play_count), 0) AS total_plays,
                    MAX(crawl_time) AS latest_crawl
                FROM video_info
                {where_sql}
            """, tuple(publish_params))
            row = cursor.fetchone()
    finally:
        db_close(db)

    try:
        r = get_redis()
        queue_length = r.llen(REDIS_START_URLS_KEY)
    except Exception:
        queue_length = 0

    return StatsResponse(
        total_videos=row['total_videos'] or 0,
        total_authors=row['total_authors'] or 0,
        total_likes=row['total_likes'] or 0,
        total_comments=row['total_comments'] or 0,
        total_shares=row['total_shares'] or 0,
        total_plays=row['total_plays'] or 0,
        latest_crawl=row['latest_crawl'],
        queue_length=queue_length,
    )


@app.get('/api/stats/authors', dependencies=[Depends(verify_read_guard)])
def stats_authors():
    db = get_db()
    try:
        with db.cursor() as cursor:
            cursor.execute("""
                SELECT COALESCE(NULLIF(TRIM(author_name), ''), '未知') AS name, COUNT(*) AS value
                FROM video_info
                GROUP BY author_name
                ORDER BY value DESC
            """)
            rows = cursor.fetchall()
    finally:
        db_close(db)
    return {'authors': rows}


@app.post('/api/crawl', response_model=CrawlResponse, dependencies=[Depends(verify_write_guard)])
def push_crawl(req: CrawlRequest):
    """只推 pending/新 id，推送成功后标记 done；Redis 不可用返回 503 且不标记。"""
    cleaned = [vid.strip() for vid in req.video_ids if vid and vid.strip()]
    records = extension_receiver.read_ids_with_status(VIDEO_IDS_PATH)
    pushable = extension_receiver.filter_pending_ids(records, cleaned)
    try:
        r = get_redis()
        count = 0
        for vid in pushable:
            task = json.dumps({
                'url': f'https://www.douyin.com/video/{vid}',
                'type': req.task_type,
            })
            r.lpush(REDIS_START_URLS_KEY, task)
            count += 1
        queue_length = r.llen(REDIS_START_URLS_KEY)
        extension_receiver.mark_ids_done(VIDEO_IDS_PATH, pushable)
    except redis.ConnectionError:
        raise HTTPException(status_code=503, detail='Redis 服务不可用')
    return CrawlResponse(
        pushed=count,
        queue_length=queue_length,
        video_ids=pushable,
        skipped=len(cleaned) - count,
    )


@app.delete('/api/videos/{video_id}', dependencies=[Depends(verify_write_guard)])
def delete_video(video_id: str):
    db = get_db()
    try:
        with db.cursor() as cursor:
            cursor.execute('DELETE FROM video_info WHERE video_id = %s', (video_id,))
            affected = cursor.rowcount
            db.commit()
        if affected == 0:
            raise HTTPException(status_code=404, detail='视频不存在')
        return {'deleted': True, 'video_id': video_id}
    finally:
        db_close(db)


@app.get('/api/queue/length', dependencies=[Depends(verify_read_guard)])
def get_queue_length():
    try:
        r = get_redis()
        return {'queue_length': r.llen(REDIS_START_URLS_KEY)}
    except redis.ConnectionError:
        raise HTTPException(status_code=503, detail='Redis 服务不可用')


@app.get('/api/queue/items', dependencies=[Depends(verify_read_guard)])
def get_queue_items(limit: int = Query(50, ge=1, le=200)):
    try:
        r = get_redis()
        raws = r.lrange(REDIS_START_URLS_KEY, 0, limit - 1)
        length = r.llen(REDIS_START_URLS_KEY)
    except redis.ConnectionError:
        raise HTTPException(status_code=503, detail='Redis 服务不可用')
    items = [it for it in (queue_service.parse_queue_item(raw) for raw in raws) if it]
    return {'queue_length': length, 'items': items}


class QueueRemoveRequest(BaseModel):
    video_ids: list[str]


@app.post('/api/queue/clear', dependencies=[Depends(verify_write_guard)])
def queue_clear():
    """清空 Redis 爬虫队列。"""
    try:
        r = get_redis()
        r.delete(REDIS_START_URLS_KEY)
    except redis.ConnectionError:
        raise HTTPException(status_code=503, detail='Redis 服务不可用')
    return {'cleared': True}


@app.post('/api/queue/remove', dependencies=[Depends(verify_write_guard)])
def queue_remove(req: QueueRemoveRequest):
    """按 video_id 批量移除队列条目（保序重建）。"""
    cleaned = [vid.strip() for vid in req.video_ids if vid and vid.strip()]
    if not cleaned:
        raise HTTPException(status_code=400, detail='没有合法的 video_id')
    try:
        r = get_redis()
        raws = r.lrange(REDIS_START_URLS_KEY, 0, -1)
        kept = queue_service.remove_items(raws, cleaned)
        removed = len(raws) - len(kept)
        if removed:
            r.delete(REDIS_START_URLS_KEY)
            if kept:
                r.rpush(REDIS_START_URLS_KEY, *kept)
        queue_length = r.llen(REDIS_START_URLS_KEY)
    except redis.ConnectionError:
        raise HTTPException(status_code=503, detail='Redis 服务不可用')
    return {'removed': removed, 'queue_length': queue_length}


# ── Spider Control ──

@app.post('/api/spider/start', dependencies=[Depends(verify_write_guard)])
def spider_start():
    ok, msg = spider_manager.start()
    if not ok:
        raise HTTPException(status_code=409, detail=msg)
    status = spider_manager.get_status()
    status['message'] = msg
    return status


@app.post('/api/spider/stop', dependencies=[Depends(verify_write_guard)])
def spider_stop():
    ok, msg = spider_manager.stop()
    if not ok:
        raise HTTPException(status_code=409, detail=msg)
    return {'message': msg, 'running': False}


@app.get('/api/spider/status', dependencies=[Depends(verify_read_guard)])
def spider_status():
    return spider_manager.get_status()


@app.get('/api/spider/log', dependencies=[Depends(verify_read_guard)])
def spider_log(lines: int = 50):
    return {'lines': spider_manager.get_log(lines)}


class QualityDeleteRequest(BaseModel):
    video_ids: list[str]


class CollectRequest(BaseModel):
    author_url: str
    max_count: int = 50


@app.post('/api/collect/author', dependencies=[Depends(verify_write_guard)])
def collect_author(req: CollectRequest):
    if not req.author_url.strip():
        raise HTTPException(status_code=400, detail='请输入作者主页链接')
    if not (1 <= req.max_count <= 100):
        raise HTTPException(status_code=400, detail='max_count 需在 1-100 之间')
    try:
        from scrapy.utils.project import get_project_settings
        cookies = get_project_settings().getdict('DOUYIN_COOKIES', {}) or {}
    except Exception:
        cookies = {}
    try:
        result = collector.collect_author_videos(
            req.author_url.strip(),
            max_count=req.max_count,
            cookies=cookies,
        )
    except collector.CollectorError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=502, detail=f'收集失败: {e}')
    return result


@app.get('/api/quality/report', dependencies=[Depends(verify_read_guard)])
def quality_report():
    db = get_db()
    try:
        with db.cursor() as cursor:
            cursor.execute('SELECT * FROM video_info')
            rows = cursor.fetchall()
    finally:
        db_close(db)
    issues = [quality_service.issue_view(r) for r in rows if quality_service.classify_row(r)]
    return {'summary': quality_service.summarize(rows), 'issues': issues}


@app.post('/api/quality/fix', dependencies=[Depends(verify_write_guard)])
def quality_fix():
    db = get_db()
    try:
        with db.cursor() as cursor:
            cursor.execute('SELECT video_id, video_title FROM video_info')
            rows = cursor.fetchall()
            fixes = quality_service.collect_title_fixes(rows)
            for video_id, title in fixes:
                cursor.execute('UPDATE video_info SET video_title = %s WHERE video_id = %s', (title, video_id))
            db.commit()
    finally:
        db_close(db)
    return {
        'fixed': len(fixes),
        'details': [{'video_id': v, 'video_title': t} for v, t in fixes[:20]],
    }


@app.post('/api/quality/delete', dependencies=[Depends(verify_write_guard)])
def quality_delete(req: QualityDeleteRequest):
    if len(req.video_ids) > quality_service.MAX_DELETE_IDS:
        raise HTTPException(status_code=400, detail=f'单次最多删除 {quality_service.MAX_DELETE_IDS} 条，请分批操作')
    db = get_db()
    try:
        deleted = 0
        rejected = []
        with db.cursor() as cursor:
            for video_id in req.video_ids:
                cursor.execute('SELECT * FROM video_info WHERE video_id = %s', (video_id,))
                row = cursor.fetchone()
                if not row:
                    rejected.append({'video_id': video_id, 'reason': '不存在'})
                    continue
                if not quality_service.is_deletable(row):
                    rejected.append({'video_id': video_id, 'reason': '当前数据已不再满足可删规则'})
                    continue
                cursor.execute('DELETE FROM video_info WHERE video_id = %s', (video_id,))
                deleted += 1
            db.commit()
    finally:
        db_close(db)
    return {'deleted': deleted, 'rejected': rejected}


@app.get('/api/quality/export', dependencies=[Depends(verify_read_guard)])
def quality_export(
    scope: str = Query('all', pattern='^(all|issues)$'),
    format: str = Query('csv', pattern='^(csv|xlsx)$'),
):
    db = get_db()
    try:
        with db.cursor() as cursor:
            cursor.execute('SELECT * FROM video_info')
            rows = cursor.fetchall()
    finally:
        db_close(db)
    if scope == 'issues':
        rows = [r for r in rows if quality_service.classify_row(r)]
    if format == 'xlsx':
        return Response(
            content=quality_service.build_xlsx(rows),
            media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={'Content-Disposition': 'attachment; filename="douyin_data.xlsx"'},
        )
    return Response(
        content=quality_service.build_csv(rows).encode('utf-8'),
        media_type='text/csv; charset=utf-8',
        headers={'Content-Disposition': 'attachment; filename="douyin_data.csv"'},
    )


class ExtensionVideosRequest(BaseModel):
    source_url: str
    videos: list[dict]


@app.post('/api/extension/videos', dependencies=[Depends(verify_write_guard)])
def extension_receive(req: ExtensionVideosRequest):
    """浏览器插件数据接收器：校验 → 批次内去重 → 部分更新 upsert。"""
    valid, rejected = extension_receiver.validate_batch(req.model_dump())
    if not valid and not rejected:
        raise HTTPException(status_code=400, detail='没有可处理的记录')
    records = extension_receiver.dedupe_records(valid)
    records, author_rejected = extension_receiver.filter_by_author_whitelist(records, ALLOWED_AUTHOR_IDS)
    rejected.extend(author_rejected)
    db = get_db()
    try:
        with db.cursor() as cursor:
            for record in records:
                sql, params = extension_receiver.build_upsert(record)
                cursor.execute(sql, params)
            db.commit()
    finally:
        db_close(db)
    return {
        'source_url': req.source_url,
        'accepted': len(valid),
        'upserted': len(records),
        'rejected': rejected,
    }


class ExtensionIdsRequest(BaseModel):
    video_ids: list[str]
    author_id: str = ''


@app.post('/api/extension/ids', dependencies=[Depends(verify_write_guard)])
def extension_save_ids(req: ExtensionIdsRequest):
    """把插件采集到的 video_id 去重追加到 video_ids.txt，供爬虫后续刷新数据。"""
    if not (1 <= len(req.video_ids) <= extension_receiver.MAX_BATCH):
        raise HTTPException(
            status_code=400,
            detail=f'video_ids 必须是 1-{extension_receiver.MAX_BATCH} 条',
        )
    cleaned: list[str] = []
    rejected: list[str] = []
    for vid in req.video_ids:
        vid = (vid or '').strip()
        if extension_receiver.validate_video_id(vid):
            cleaned.append(vid)
        else:
            rejected.append(vid)
    if not cleaned:
        raise HTTPException(status_code=400, detail='没有合法的 video_id')
    added, total = extension_receiver.append_ids_file(
        VIDEO_IDS_PATH, cleaned, author_id=(req.author_id or '').strip(),
    )
    return {'added': added, 'total': total, 'rejected': rejected}


@app.get('/api/extension/ids', dependencies=[Depends(verify_read_guard)])
def extension_list_ids():
    """返回 video_ids.txt 的数量、纯 id 列表与带状态/作者/昵称明细，供前端查看/导入爬虫队列。"""
    records = extension_receiver.read_ids_with_status(VIDEO_IDS_PATH)
    author_ids = {r['author_id'] for r in records if r['author_id']}
    author_map: dict = {}
    if author_ids:
        db = get_db()
        try:
            with db.cursor() as cursor:
                placeholders = ', '.join(['%s'] * len(author_ids))
                cursor.execute(
                    f'SELECT DISTINCT author_id, author_name FROM video_info '
                    f'WHERE author_id IN ({placeholders})',
                    tuple(author_ids),
                )
                for row in cursor.fetchall():
                    author_map[row['author_id']] = row['author_name'] or ''
        finally:
            db_close(db)
    items = extension_receiver.attach_author_names(records, author_map)
    return {
        'total': len(records),
        'video_ids': [r['video_id'] for r in records],
        'items': items,
    }


@app.put('/api/extension/ids', dependencies=[Depends(verify_write_guard)])
def extension_replace_ids(req: ExtensionIdsRequest):
    """前端直接编辑保存/清空：覆盖写入 video_ids.txt（空列表=清空，锁 + 原子替换）。"""
    if len(req.video_ids) > 2000:
        raise HTTPException(status_code=400, detail='video_ids 数量超限（最多 2000 条）')
    cleaned: list[str] = []
    rejected: list[str] = []
    for vid in req.video_ids:
        vid = (vid or '').strip()
        if extension_receiver.validate_video_id(vid):
            cleaned.append(vid)
        else:
            rejected.append(vid)
    total = extension_receiver.write_ids_file(VIDEO_IDS_PATH, cleaned)
    return {'total': total, 'rejected': rejected}


class ExtensionIdsStatusRequest(BaseModel):
    video_ids: list[str]
    status: str


@app.post('/api/extension/ids/status', dependencies=[Depends(verify_write_guard)])
def extension_set_ids_status(req: ExtensionIdsStatusRequest):
    """批量切换 id 状态（pending/done），供前端强制重爬/标记。"""
    if req.status not in ('pending', 'done'):
        raise HTTPException(status_code=400, detail='status 必须是 pending 或 done')
    cleaned: list[str] = []
    rejected: list[str] = []
    for vid in req.video_ids:
        vid = (vid or '').strip()
        if extension_receiver.validate_video_id(vid):
            cleaned.append(vid)
        else:
            rejected.append(vid)
    if not cleaned:
        raise HTTPException(status_code=400, detail='没有合法的 video_id')
    updated = extension_receiver.set_ids_status(VIDEO_IDS_PATH, cleaned, req.status)
    return {'updated': updated, 'rejected': rejected}


@app.get('/api/analyze/authors', dependencies=[Depends(verify_read_guard)])
def analyze_authors():
    """作者下拉数据源：author_id + author_name + 视频数。"""
    db = get_db()
    try:
        with db.cursor() as cursor:
            author_clause, author_params = extension_receiver.build_author_filter(ALLOWED_AUTHOR_IDS)
            where_sql = 'author_id IS NOT NULL AND author_id <> \'\''
            params = []
            if author_clause:
                where_sql += ' AND ' + author_clause
                params.extend(author_params)
            cursor.execute(f"""
                SELECT author_id, author_name, COUNT(*) AS count
                FROM video_info
                WHERE {where_sql}
                GROUP BY author_id, author_name
                ORDER BY count DESC
            """, tuple(params))
            rows = cursor.fetchall()
    finally:
        db_close(db)
    return {'authors': rows}


@app.get('/api/analyze/personal', dependencies=[Depends(verify_read_guard)])
def analyze_personal(
    author_id: str = Query(..., description='作者 uid'),
    sort_by: str = Query('likes', description='Top 视频排序维度'),
    start_date: str = Query('', description='发布时间起始（YYYY-MM-DD）'),
    end_date: str = Query('', description='发布时间结束（YYYY-MM-DD）'),
):
    """按作者聚合个人分析：概览 / 发布趋势 / 播放趋势 / Top 视频。"""
    if sort_by not in ('likes', 'plays', 'comments', 'shares', 'collects', 'engagement'):
        raise HTTPException(status_code=400, detail='sort_by 必须是 likes/plays/comments/shares/collects/engagement')
    db = get_db()
    try:
        with db.cursor() as cursor:
            publish_clause, publish_params = apply_publish_filter(start_date, end_date)
            author_clause, author_params = extension_receiver.build_author_filter(ALLOWED_AUTHOR_IDS)
            where_sql = 'author_id = %s'
            where_params = [author_id]
            if publish_clause:
                where_sql += ' AND ' + publish_clause
                where_params.extend(publish_params)
            if author_clause:
                where_sql += ' AND ' + author_clause
                where_params.extend(author_params)
            cursor.execute(f'SELECT * FROM video_info WHERE {where_sql}', tuple(where_params))
            rows = cursor.fetchall()
    finally:
        db_close(db)
    author_name = (rows[0].get('author_name') or '') if rows else ''
    return {
        'author_id': author_id,
        'author_name': author_name,
        'summary': analyzer.summarize_rows(rows),
        'trend': analyzer.build_trend(rows),
        'play_trend': analyzer.build_play_trend(rows),
        'top_videos': analyzer.top_videos(rows, sort_by=sort_by),
    }


@app.get('/api/export', dependencies=[Depends(verify_read_guard)])
def export_data(
    search: str = Query('', description='搜索视频标题/作者/ID'),
    sort_by: str = Query('crawl_time', description='排序字段'),
    order: str = Query('desc', pattern='^(asc|desc)$'),
    start_date: str = Query('', description='发布时间起始（YYYY-MM-DD）'),
    end_date: str = Query('', description='发布时间结束（YYYY-MM-DD）'),
    format: str = Query('csv', pattern='^(csv|xlsx)$'),
):
    """导出当前筛选结果（与 /api/videos 同参数，不分页，上限 10000）。"""
    allowed_sort = {
        'video_id', 'video_title', 'author_name', 'publish_time',
        'like_count', 'comment_count', 'share_count', 'play_count', 'collect_count',
        'crawl_time', 'update_time',
    }
    if sort_by not in allowed_sort:
        sort_by = 'crawl_time'
    order_clause = 'DESC' if order == 'desc' else 'ASC'
    publish_clause, publish_params = apply_publish_filter(start_date, end_date)
    where_parts = []
    params = []
    if search:
        where_parts.append('(video_id LIKE %s OR video_title LIKE %s OR author_name LIKE %s)')
        params.extend([f'%{search}%'] * 3)
    if publish_clause:
        where_parts.append(publish_clause)
        params.extend(publish_params)
    where_sql = ('WHERE ' + ' AND '.join(where_parts)) if where_parts else ''

    db = get_db()
    try:
        with db.cursor() as cursor:
            cursor.execute(f'SELECT COUNT(*) AS n FROM video_info {where_sql}', tuple(params))
            total = cursor.fetchone()['n']
        _check_export_total(total)
        if format == 'csv':
            def gen():
                conn = get_db()
                try:
                    with conn.cursor(pymysql.cursors.SSCursor) as cursor:
                        cursor.execute(
                            f'SELECT * FROM video_info {where_sql} ORDER BY {sort_by} {order_clause}',
                            tuple(params),
                        )
                        buf = io.StringIO()
                        buf.write('\ufeff')
                        writer = csv.DictWriter(buf, fieldnames=export_service.EXPORT_COLUMNS, extrasaction='ignore')
                        writer.writeheader()
                        yield buf.getvalue()
                        while True:
                            batch = cursor.fetchmany(1000)
                            if not batch:
                                break
                            buf = io.StringIO()
                            writer = csv.DictWriter(buf, fieldnames=export_service.EXPORT_COLUMNS, extrasaction='ignore')
                            for row in batch:
                                writer.writerow({c: ('' if row.get(c) is None else row.get(c)) for c in export_service.EXPORT_COLUMNS})
                            yield buf.getvalue()
                finally:
                    db_close(conn)
            return StreamingResponse(
                gen(), media_type='text/csv; charset=utf-8',
                headers={'Content-Disposition': 'attachment; filename="douyin_data.csv"'},
            )
        tmp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        tmp.close()
        from openpyxl import Workbook
        wb = Workbook(write_only=True)
        ws = wb.create_sheet()
        ws.append(list(export_service.EXPORT_COLUMNS))
        conn = get_db()
        try:
            with conn.cursor(pymysql.cursors.SSCursor) as cursor:
                cursor.execute(
                    f'SELECT * FROM video_info {where_sql} ORDER BY {sort_by} {order_clause}',
                    tuple(params),
                )
                while True:
                    batch = cursor.fetchmany(1000)
                    if not batch:
                        break
                    for row in batch:
                        ws.append([('' if row.get(c) is None else row.get(c)) for c in export_service.EXPORT_COLUMNS])
        finally:
            db_close(conn)
        wb.save(tmp.name)
        return FileResponse(
            tmp.name,
            media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            filename='douyin_data.xlsx',
        )
    finally:
        db_close(db)


class CleanupToggleRequest(BaseModel):
    enabled: bool


@app.get('/api/cleanup/status', dependencies=[Depends(verify_read_guard)])
def cleanup_status():
    """返回定时清理开关、上次执行时间、条数与指定作者。"""
    cfg = _read_cleanup_config()
    return {
        'enabled': bool(cfg['enabled']),
        'last_clean_time': cfg['last_clean_time'],
        'batch_size': int(cfg['batch_size']),
        'authors': list(cfg['authors']),
    }


@app.post('/api/cleanup/toggle', dependencies=[Depends(verify_write_guard)])
def cleanup_toggle(req: CleanupToggleRequest):
    """切换定时清理开关（写接口，走令牌守卫）。"""
    cfg = _read_cleanup_config()
    cfg['enabled'] = req.enabled
    _write_cleanup_config(cfg)
    return {'enabled': req.enabled}


class CleanupSettingsRequest(BaseModel):
    batch_size: int = 200
    authors: list[str] = []


@app.post('/api/cleanup/settings', dependencies=[Depends(verify_write_guard)])
def cleanup_settings(req: CleanupSettingsRequest):
    """保存清理条数与指定作者（空 authors = 全部作者）。"""
    if not (1 <= req.batch_size <= 1000):
        raise HTTPException(status_code=400, detail='batch_size 必须在 1-1000 之间')
    cfg = _read_cleanup_config()
    cfg['batch_size'] = req.batch_size
    cfg['authors'] = list(req.authors)
    _write_cleanup_config(cfg)
    return {'batch_size': req.batch_size, 'authors': req.authors}


FRONTEND_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'frontend')
DIST_DIR = os.path.join(FRONTEND_DIR, 'dist')

if os.path.isdir(DIST_DIR):
    # Vue 应用：/app 子路径部署，history 路由刷新由 SPA 兜底
    app.mount('/app/assets', StaticFiles(directory=os.path.join(DIST_DIR, 'assets')), name='frontend-assets')

    @app.get('/app')
    @app.get('/app/{full_path:path}')
    def frontend_spa(full_path: str = ''):
        index = os.path.join(DIST_DIR, 'index.html')
        if full_path:
            target = os.path.abspath(os.path.join(DIST_DIR, full_path))
            if target.startswith(os.path.abspath(DIST_DIR)) and os.path.isfile(target):
                return FileResponse(target)
        return FileResponse(index)



@app.get('/', include_in_schema=False)
def root_redirect():
    """根路径重定向到 Vue 应用。"""
    return RedirectResponse('/app/')


if __name__ == '__main__':
    import uvicorn
    uvicorn.run('api:app', host='0.0.0.0', port=8001, reload=False)
